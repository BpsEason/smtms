"""Parse, validate, and import supplier master data with login accounts from Excel.

This module keeps the Excel layout, validation rules, and the write path in one
place so the API layer stays thin. Passwords are only ever returned to the caller
once (as the initial password); they are never persisted in plaintext.
"""

from __future__ import annotations

import re
import secrets
from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.security import get_password_hash
from app.models.quality import SupplierORM
from app.models.user import PermissionORM, RoleORM, RolePermissionORM, UserORM, UserRoleORM

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000

SHEET_TITLE = "suppliers"

TEMPLATE_HEADERS: tuple[str, ...] = (
    "supplier_code",
    "supplier_name",
    "tax_id",
    "contact_name",
    "phone",
    "email",
    "country",
    "address",
    "status",
    "login_username",
    "login_email",
    "login_password",
)

REQUIRED_HEADERS: tuple[str, ...] = ("supplier_code", "supplier_name", "login_username")

HEADER_EXAMPLE_ROW: tuple[str, ...] = (
    "SUP-1001",
    "範例供應商股份有限公司",
    "12345678",
    "王小明",
    "02-1234-5678",
    "contact@example.com",
    "Taiwan",
    "台北市信義區範例路 1 號",
    "active",
    "sup1001",
    "contact@example.com",
    "",
)

SUPPLIER_ROLE_NAME = "Supplier Portal User"
SUPPLIER_ROLE_DESCRIPTION = "供應商入口帳號（僅供供應商登入使用，不得授予內部管理權限）"
SUPPLIER_ROLE_PERMISSION = "suppliers:read"

_USERNAME_PATTERN = re.compile(r"^[A-Za-z0-9._-]{1,80}$")
_EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_PASSWORD_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZabcdefghijkmnopqrstuvwxyz23456789"
_DEFAULT_PASSWORD_LENGTH = 12


class ImportFormatError(Exception):
    """Raised when the uploaded workbook cannot be read as a supplier import file."""

    def __init__(self, errors: list["RowError"]) -> None:
        super().__init__("Supplier import file is not usable")
        self.errors = errors


@dataclass(frozen=True)
class RowError:
    sheet: str
    row: int
    field: str
    message: str


@dataclass
class ParsedRow:
    sheet: str
    row: int
    values: dict[str, str | None]


@dataclass
class ParsedWorkbook:
    sheet: str
    rows: list[ParsedRow]


@dataclass
class ValidatedRow:
    sheet: str
    row: int
    supplier_code: str
    supplier_name: str
    tax_id: str | None
    contact_name: str | None
    phone: str | None
    email: str | None
    country: str | None
    address: str | None
    status: str
    login_username: str
    login_email: str | None
    login_password: str | None
    skip: bool


@dataclass
class ValidationOutcome:
    sheet: str
    total_rows: int
    rows: list[ValidatedRow]
    errors: list[RowError]

    @property
    def skipped_rows(self) -> list[ValidatedRow]:
        return [row for row in self.rows if row.skip]


@dataclass
class CreatedAccount:
    sheet: str
    row: int
    supplier_id: int
    supplier_code: str
    login_username: str
    initial_password: str
    password_generated: bool


def build_template_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None, "Worksheet should always exist for new workbook"
    sheet.title = SHEET_TITLE
    sheet.append(list(TEMPLATE_HEADERS))
    sheet.append(list(HEADER_EXAMPLE_ROW))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index in range(1, len(TEMPLATE_HEADERS) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 22
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _clean_cell(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _map_headers(header_row: tuple[object, ...]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, raw in enumerate(header_row):
        if raw is None:
            continue
        key = str(raw).strip().lower()
        if key in TEMPLATE_HEADERS and key not in mapping:
            mapping[key] = index
    return mapping


def parse_supplier_workbook(content: bytes) -> ParsedWorkbook:
    if not content:
        raise ImportFormatError([RowError("", 0, "file", "檔案為空，請重新選擇 Excel 檔案。")])
    if len(content) > MAX_FILE_BYTES:
        raise ImportFormatError([RowError("", 0, "file", "檔案超過大小上限（5 MB）。")])

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated exception types.
        raise ImportFormatError(
            [RowError("", 0, "file", "無法解析 Excel 檔案，請確認為 .xlsx 格式且未損毀。")]
        ) from exc

    try:
        sheet = workbook.active
        if sheet is None:
            raise ImportFormatError([RowError("", 0, "file", "找不到工作表。")])
        assert sheet is not None
        sheet_name = sheet.title
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as exc:
            raise ImportFormatError(
                [RowError(sheet_name, 1, "header", "工作表沒有內容。")]
            ) from exc

        header_map = _map_headers(tuple(header_row))
        missing = [header for header in REQUIRED_HEADERS if header not in header_map]
        if missing:
            raise ImportFormatError(
                [
                    RowError(
                        sheet_name,
                        1,
                        ",".join(missing),
                        f"工作表缺少必要欄位：{', '.join(missing)}。",
                    )
                ]
            )

        parsed_rows: list[ParsedRow] = []
        for index, raw in enumerate(rows, start=2):
            if len(parsed_rows) >= MAX_ROWS:
                raise ImportFormatError(
                    [RowError(sheet_name, index, "rows", f"資料筆數超過上限（{MAX_ROWS} 筆）。")]
                )
            values = {
                header: _clean_cell(raw[column] if column < len(raw) else None)
                for header, column in header_map.items()
            }
            if all(value is None for value in values.values()):
                continue
            parsed_rows.append(ParsedRow(sheet=sheet_name, row=index, values=values))

        if not parsed_rows:
            raise ImportFormatError([RowError(sheet_name, 1, "rows", "查無可匯入的資料列。")])
        return ParsedWorkbook(sheet=sheet_name, rows=parsed_rows)
    finally:
        workbook.close()


def _generate_password() -> str:
    return "".join(secrets.choice(_PASSWORD_ALPHABET) for _ in range(_DEFAULT_PASSWORD_LENGTH))


def validate_supplier_rows(db: Session, parsed: ParsedWorkbook) -> ValidationOutcome:
    sheet = parsed.sheet
    existing_suppliers = {
        supplier.supplier_code: supplier for supplier in db.scalars(select(SupplierORM)).all()
    }
    existing_tax_ids = {
        supplier.tax_id: supplier
        for supplier in existing_suppliers.values()
        if supplier.tax_id
    }
    existing_usernames = {user.username for user in db.scalars(select(UserORM)).all()}
    linked_usernames = {
        user.supplier_id: user.username
        for user in db.scalars(select(UserORM).where(UserORM.supplier_id.is_not(None))).all()
    }

    seen_codes: dict[str, int] = {}
    seen_tax_ids: dict[str, int] = {}
    seen_usernames: dict[str, int] = {}
    errors: list[RowError] = []
    validated: list[ValidatedRow] = []

    def add_error(row: int, column: str, message: str) -> None:
        errors.append(RowError(sheet=sheet, row=row, field=column, message=message))

    for item in parsed.rows:
        row = item.row
        values = item.values
        row_errors_before = len(errors)

        supplier_code = values.get("supplier_code")
        supplier_name = values.get("supplier_name")
        tax_id = values.get("tax_id")
        contact_name = values.get("contact_name")
        phone = values.get("phone")
        email = values.get("email")
        country = values.get("country")
        address = values.get("address")
        status = values.get("status") or "active"
        login_username = values.get("login_username")
        login_email = values.get("login_email")
        login_password = values.get("login_password")

        for column, value in (
            ("supplier_code", supplier_code),
            ("supplier_name", supplier_name),
            ("login_username", login_username),
        ):
            if not value:
                add_error(row, column, "此欄位為必填。")

        for column, value, limit in (
            ("supplier_code", supplier_code, 80),
            ("supplier_name", supplier_name, 200),
            ("tax_id", tax_id, 20),
            ("contact_name", contact_name, 120),
            ("phone", phone, 40),
            ("email", email, 160),
            ("country", country, 80),
            ("address", address, 255),
            ("status", status, 20),
            ("login_email", login_email, 255),
        ):
            if value and len(value) > limit:
                add_error(row, column, f"長度不可超過 {limit} 個字元。")

        if login_username and not _USERNAME_PATTERN.match(login_username):
            add_error(row, "login_username", "僅允許英文字母、數字、句點、底線與連字號。")
        if email and not _EMAIL_PATTERN.match(email):
            add_error(row, "email", "電子郵件格式不正確。")
        if login_email and not _EMAIL_PATTERN.match(login_email):
            add_error(row, "login_email", "電子郵件格式不正確。")
        if login_password and not 8 <= len(login_password.encode("utf-8")) <= 72:
            add_error(row, "login_password", "密碼長度需介於 8 至 72 個位元組。")

        if supplier_code:
            if supplier_code in seen_codes:
                add_error(
                    row,
                    "supplier_code",
                    f"與第 {seen_codes[supplier_code]} 列的供應商編號重複。",
                )
            else:
                seen_codes[supplier_code] = row
        if tax_id:
            if tax_id in seen_tax_ids:
                add_error(row, "tax_id", f"與第 {seen_tax_ids[tax_id]} 列的統一編號重複。")
            else:
                seen_tax_ids[tax_id] = row
        if login_username:
            if login_username in seen_usernames:
                add_error(
                    row,
                    "login_username",
                    f"與第 {seen_usernames[login_username]} 列的登入帳號重複。",
                )
            else:
                seen_usernames[login_username] = row

        skip = False
        if supplier_code and supplier_code in existing_suppliers:
            existing = existing_suppliers[supplier_code]
            if login_username and linked_usernames.get(existing.id) == login_username:
                skip = True
            else:
                add_error(row, "supplier_code", "供應商編號已存在，且無法對應既有的登入帳號。")
        if not skip:
            if tax_id and tax_id in existing_tax_ids:
                add_error(row, "tax_id", "統一編號已存在於系統中。")
            if login_username and login_username in existing_usernames:
                add_error(row, "login_username", "登入帳號已存在於系統中。")

        if len(errors) != row_errors_before:
            continue

        validated.append(
            ValidatedRow(
                sheet=sheet,
                row=row,
                supplier_code=supplier_code or "",
                supplier_name=supplier_name or "",
                tax_id=tax_id,
                contact_name=contact_name,
                phone=phone,
                email=email,
                country=country,
                address=address,
                status=status,
                login_username=login_username or "",
                login_email=login_email,
                login_password=login_password,
                skip=skip,
            )
        )

    return ValidationOutcome(
        sheet=sheet,
        total_rows=len(parsed.rows),
        rows=validated,
        errors=errors,
    )


def _ensure_supplier_role(db: Session) -> RoleORM:
    role = db.scalar(select(RoleORM).where(RoleORM.name == SUPPLIER_ROLE_NAME))
    if role is None:
        role = RoleORM(
            name=SUPPLIER_ROLE_NAME,
            description=SUPPLIER_ROLE_DESCRIPTION,
        )
        db.add(role)
        db.flush()

    permission = db.scalar(
        select(PermissionORM).where(PermissionORM.code == SUPPLIER_ROLE_PERMISSION)
    )
    if permission is None:
        permission = PermissionORM(
            code=SUPPLIER_ROLE_PERMISSION,
            name="讀取供應商資料",
            description="供應商入口帳號可讀取供應商基本資料。",
        )
        db.add(permission)
        db.flush()

    link = db.scalar(
        select(RolePermissionORM.id).where(
            RolePermissionORM.role_id == role.id,
            RolePermissionORM.permission_id == permission.id,
        )
    )
    if link is None:
        db.add(
            RolePermissionORM(
                role_id=role.id,
                permission_id=permission.id,
            )
        )
        db.flush()
    return role


def execute_supplier_import(db: Session, outcome: ValidationOutcome) -> list[CreatedAccount]:
    """Create suppliers and their login accounts, or create nothing at all.

    Callers must ensure ``outcome.errors`` is empty. All rows share one
    transaction so a supplier is never created without its account.
    """
    rows_to_create = [row for row in outcome.rows if not row.skip]
    if not rows_to_create:
        return []

    role = _ensure_supplier_role(db)

    created: list[CreatedAccount] = []
    for row in rows_to_create:
        supplier = SupplierORM(
            supplier_code=row.supplier_code,
            supplier_name=row.supplier_name,
            tax_id=row.tax_id,
            contact_name=row.contact_name,
            phone=row.phone,
            email=row.email,
            status=row.status,
            address=row.address,
            country=row.country,
        )
        db.add(supplier)

        password_generated = row.login_password is None
        password = row.login_password or _generate_password()
        user = UserORM(
            username=row.login_username,
            email=row.login_email,
            password_hash=get_password_hash(password),
            is_active=True,
            supplier_id=None,
            must_change_password=True,
        )
        db.add(user)
        db.flush()
        user.supplier_id = supplier.id

        db.add(
            UserRoleORM(
                user_id=user.id,
                role_id=role.id,
            )
        )

        created.append(
            CreatedAccount(
                sheet=row.sheet,
                row=row.row,
                supplier_id=supplier.id,
                supplier_code=row.supplier_code,
                login_username=row.login_username,
                initial_password=password,
                password_generated=password_generated,
            )
        )

    db.commit()
    return created